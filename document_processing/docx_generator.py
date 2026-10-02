"""DOCX export using python-docx."""

from __future__ import annotations

import logging
from io import BytesIO
from pathlib import Path
from typing import List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from backend.config import get_settings
from backend.utils.sanitizers import sanitize_text
from document_processing.formatter import DISCLAIMER, build_footer_text, split_into_blocks

logger = logging.getLogger(__name__)


def _set_run_font(run, name: str = "Times New Roman", size: int = 11) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)


def _add_page_number(paragraph) -> None:
    """Add PAGE field to a paragraph for page numbering."""
    run = paragraph.add_run()
    fld_char_begin = run._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    run._element.append(fld_char_begin)

    run2 = paragraph.add_run()
    instr = run2._element.makeelement(qn("w:instrText"), {})
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run2._element.append(instr)

    run3 = paragraph.add_run()
    fld_char_end = run3._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    run3._element.append(fld_char_end)


def generate_docx(
    title: str,
    content: str,
    parties: Optional[List[str]] = None,
    effective_date: Optional[str] = None,
    document_type: Optional[str] = None,
    include_logo: bool = True,
    include_signature: bool = True,
) -> bytes:
    """Generate a DOCX document and return raw bytes."""
    parties = parties or []
    title = sanitize_text(title, max_length=300)
    content = sanitize_text(content)
    settings = get_settings()

    doc = Document()

    # Default style
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.15

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # Logo
    if include_logo:
        logo = settings.logo_path
        if logo.exists():
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str(logo), width=Inches(1.2))
            except Exception:  # noqa: BLE001
                logger.warning("Could not embed logo into DOCX")

    # Title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run(title.upper())
    _set_run_font(run, size=16)
    run.bold = True

    if document_type:
        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = meta.add_run(document_type)
        _set_run_font(run, size=11)
        run.italic = True

    # Parties / date meta
    if parties:
        p = doc.add_paragraph()
        run = p.add_run("Parties: ")
        _set_run_font(run, size=11)
        run.bold = True
        run2 = p.add_run(", ".join(parties))
        _set_run_font(run2, size=11)

    if effective_date:
        p = doc.add_paragraph()
        run = p.add_run("Effective Date: ")
        _set_run_font(run, size=11)
        run.bold = True
        run2 = p.add_run(effective_date)
        _set_run_font(run2, size=11)

    doc.add_paragraph()

    # Body blocks
    for block in split_into_blocks(content):
        lines = block.splitlines()
        first = lines[0].strip() if lines else ""
        is_heading = bool(
            first
            and len(first) < 90
            and (
                first.isupper()
                or first.endswith(":")
                or first.startswith("#")
            )
        )
        if is_heading:
            heading_text = first.lstrip("#").strip().rstrip(":")
            hp = doc.add_paragraph()
            run = hp.add_run(heading_text)
            _set_run_font(run, size=12)
            run.bold = True
            body_lines = lines[1:]
        else:
            body_lines = lines

        for line in body_lines:
            stripped = line.strip()
            if not stripped:
                continue
            if stripped.startswith(("- ", "* ", "• ")):
                bp = doc.add_paragraph(style="List Bullet")
                run = bp.add_run(stripped[2:].strip())
                _set_run_font(run, size=11)
            else:
                para = doc.add_paragraph()
                run = para.add_run(stripped)
                _set_run_font(run, size=11)

    # Optional simple terms table when pipe-separated rows exist
    _maybe_add_terms_table(doc, content)

    if include_signature and "signature" not in content.lower():
        doc.add_paragraph()
        hp = doc.add_paragraph()
        run = hp.add_run("SIGNATURES")
        _set_run_font(run, size=12)
        run.bold = True
        for party in parties or ["Party A", "Party B"]:
            doc.add_paragraph()
            p = doc.add_paragraph()
            run = p.add_run(f"For {party}:")
            _set_run_font(run, size=11)
            run.bold = True
            doc.add_paragraph("Signature: _______________________________")
            doc.add_paragraph("Name: ____________________________________")
            doc.add_paragraph("Date: ____________________________________")

    # Disclaimer
    doc.add_paragraph()
    disc = doc.add_paragraph()
    run = disc.add_run(DISCLAIMER)
    _set_run_font(run, size=9)
    run.italic = True
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # Footer with page number
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run(build_footer_text() + "  |  Page ")
    _set_run_font(run, size=8)
    _add_page_number(fp)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _maybe_add_terms_table(doc: Document, content: str) -> None:
    """If content contains simple 'Term | Value' lines, render a table."""
    rows = []
    for line in content.splitlines():
        if "|" in line and not line.strip().startswith("#"):
            parts = [p.strip() for p in line.split("|")]
            if len(parts) == 2 and parts[0] and parts[1]:
                rows.append((parts[0], parts[1]))
    if len(rows) < 2:
        return
    doc.add_paragraph()
    hp = doc.add_paragraph()
    run = hp.add_run("Terms Summary")
    _set_run_font(run, size=12)
    run.bold = True
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Term"
    hdr[1].text = "Value"
    for term, value in rows[:30]:
        row = table.add_row().cells
        row[0].text = term
        row[1].text = value
